"""Tests for issue 10: student import from Excel (importar-aluno).

Covers the functional specification:
- labels are found by name (not by row number) in the first column of every
  sheet, tolerating different row positions between files ("treino" starts at
  row 2, "dieta" at row 1);
- "atleta" becomes the name with repeated inner spaces collapsed and no
  recapitalization/spell-fixing;
- "objetivo" is imported as written, with no correction;
- "periodo" extracts DD/MM/AAAA dates by regex, tolerating garbage between
  them; a single date becomes plan_start only, with a warning;
- "idade" is never imported and is reported in not_imported with its raw
  value quoted;
- future slices (training plan, diet, macro/mesocycle, observations) are
  always reported as not imported;
- the original .xlsx files are opened read-only and never modified;
- a folder without any .xlsx, or a folder that does not exist, raises
  FileNotFoundError with a clear Portuguese message;
- the CLI ("importar-aluno") prints the report, writes the student to the
  database, refuses to duplicate an existing name, supports --simular
  (report only, no write), and fails cleanly (exit code 1, no traceback)
  when the name is missing or the folder does not exist.

Isolation follows the same pattern as tests/test_lista_alunos.py:
KAIROS_DATA_DIR points at a temp dir and the cached engine is disposed on
teardown. The Excel-parsing tests (parse_aluno_folder) never touch the
database at all.
"""

import sqlite3
from pathlib import Path
from typing import Optional

import pytest
from openpyxl import Workbook

import kairos.db
from kairos.alunos.excel_import import parse_aluno_folder
from kairos.alunos.service import find_aluno_by_name
from kairos.cli import main
from kairos.config import db_path

# Garbage character the coach's spreadsheet software sometimes leaves between
# the two dates of the "periodo" cell (mirrors the real file: "�").
GARBLED_PERIOD = "inicio :02/01/2024 �termino 15/02/2024"

REAL_ALUNO_FOLDER = Path(
    r"C:\Users\jeffr\OneDrive\Ambiente de Trabalho\consultoria\alunos\marcos anastasio"
)


def _build_treino_workbook(path: Path, period_text: str = GARBLED_PERIOD) -> None:
    """Build a workbook imitating the real "treino" file: labels start at row 2."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Treino"
    sheet["A1"] = "Ficha de treino"
    sheet["A2"] = "atleta"
    sheet["B2"] = "Marcos  anastasio"
    sheet["A3"] = "idade"
    sheet["B3"] = "29 anos"
    sheet["A4"] = "periodo"
    sheet["B4"] = period_text
    sheet["A5"] = "objetivo"
    sheet["B5"] = "definiçao e hipertrofia"
    # Loose rows below simulating the weekly training split (not recognized labels).
    sheet["A7"] = "segunda"
    sheet["B7"] = "peito e triceps"
    sheet["A8"] = "terca"
    sheet["B8"] = "costas e biceps"
    workbook.save(path)


def _build_dieta_workbook(path: Path, period_text: str = GARBLED_PERIOD) -> None:
    """Build a workbook imitating the real "dieta" file: labels start at row 1."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dieta"
    sheet["A1"] = "atleta"
    sheet["B1"] = "Marcos  anastasio"
    sheet["A2"] = "idade"
    sheet["B2"] = "29 anos"
    sheet["A3"] = "periodo"
    sheet["B3"] = period_text
    sheet["A4"] = "objetivo"
    sheet["B4"] = "definiçao e hipertrofia"
    # Loose row below simulating a diet meal (not a recognized label).
    sheet["A6"] = "café da manhã"
    sheet["B6"] = "ovos e aveia"
    workbook.save(path)


@pytest.fixture
def aluno_folder(tmp_path: Path) -> Path:
    """A student folder with the two real-world spreadsheets (treino + dieta)."""
    folder = tmp_path / "marcos anastasio"
    folder.mkdir()
    _build_treino_workbook(folder / "treino.xlsx")
    _build_dieta_workbook(folder / "dieta.xlsx")
    return folder


@pytest.fixture
def data_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """Point KAIROS_DATA_DIR at a temp dir and clean up the engine cache."""
    target = tmp_path / "kairos-data"
    monkeypatch.setenv("KAIROS_DATA_DIR", str(target))
    kairos.db.dispose_engine()  # drop any engine cached for a previous URL
    yield target
    kairos.db.dispose_engine()  # release the SQLite file handle (Windows)


def _count_alunos() -> int:
    """Count rows in the "alunos" table using a raw SQLite connection."""
    connection = sqlite3.connect(db_path())
    try:
        cursor = connection.execute("SELECT COUNT(*) FROM alunos")
        return cursor.fetchone()[0]
    finally:
        connection.close()


def _fetch_aluno_row() -> Optional[sqlite3.Row]:
    """Fetch the single "alunos" row as a dict-like sqlite3.Row, or None."""
    connection = sqlite3.connect(db_path())
    connection.row_factory = sqlite3.Row
    try:
        cursor = connection.execute(
            "SELECT name, objective, plan_start, plan_end FROM alunos"
        )
        return cursor.fetchone()
    finally:
        connection.close()


# --- parser tests (no database involved) -----------------------------------


def test_parser_recognizes_name_objective_and_period(aluno_folder: Path) -> None:
    parsed = parse_aluno_folder(aluno_folder)

    assert parsed.fields["name"] == "Marcos anastasio"
    assert parsed.fields["objective"] == "definiçao e hipertrofia"
    assert parsed.fields["plan_start"] == "2024-01-02"
    assert parsed.fields["plan_end"] == "2024-02-15"


def test_parser_never_imports_age_and_reports_it(aluno_folder: Path) -> None:
    parsed = parse_aluno_folder(aluno_folder)

    assert "age" not in parsed.fields
    assert "idade" not in parsed.fields
    items = [item for item, _reason in parsed.not_imported]
    assert 'idade "29 anos"' in items


def test_parser_reports_future_slices_training_and_diet(aluno_folder: Path) -> None:
    parsed = parse_aluno_folder(aluno_folder)

    items = [item for item, _reason in parsed.not_imported]
    assert any("treino" in item for item in items)
    assert any("dieta" in item for item in items)


def test_parser_finds_labels_regardless_of_row_position(aluno_folder: Path) -> None:
    # "treino.xlsx" has its labels starting at row 2, "dieta.xlsx" at row 1;
    # both files being fully read (and the name/objective/period being
    # correctly recognized, as asserted above) proves the label search is by
    # name and not by row index.
    parsed = parse_aluno_folder(aluno_folder)

    assert len(parsed.source_files) == 2
    assert set(parsed.source_files) == {"treino.xlsx", "dieta.xlsx"}


def test_parser_single_date_period_fills_only_start_and_warns(
    tmp_path: Path,
) -> None:
    folder = tmp_path / "single_date"
    folder.mkdir()
    _build_treino_workbook(folder / "treino.xlsx", period_text="inicio: 02/01/2024")
    _build_dieta_workbook(folder / "dieta.xlsx", period_text="inicio: 02/01/2024")

    parsed = parse_aluno_folder(folder)

    assert parsed.fields["plan_start"] == "2024-01-02"
    assert "plan_end" not in parsed.fields
    assert any("Apenas uma data" in warning for warning in parsed.warnings)


def test_parser_folder_without_xlsx_raises_file_not_found(tmp_path: Path) -> None:
    empty_folder = tmp_path / "vazio"
    empty_folder.mkdir()

    with pytest.raises(FileNotFoundError, match="Nenhum arquivo .xlsx encontrado"):
        parse_aluno_folder(empty_folder)


def test_parser_nonexistent_folder_raises_file_not_found(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        parse_aluno_folder(tmp_path / "does-not-exist")


def test_parser_never_modifies_original_files(aluno_folder: Path) -> None:
    treino_path = aluno_folder / "treino.xlsx"
    dieta_path = aluno_folder / "dieta.xlsx"

    treino_before = treino_path.stat()
    dieta_before = dieta_path.stat()

    parse_aluno_folder(aluno_folder)

    treino_after = treino_path.stat()
    dieta_after = dieta_path.stat()

    assert treino_before.st_mtime == treino_after.st_mtime
    assert treino_before.st_size == treino_after.st_size
    assert dieta_before.st_mtime == dieta_after.st_mtime
    assert dieta_before.st_size == dieta_after.st_size


# --- CLI tests ---------------------------------------------------------------


def test_cli_simular_prints_report_and_does_not_write(
    data_dir: Path, aluno_folder: Path, capsys: pytest.CaptureFixture
) -> None:
    exit_code = main(["importar-aluno", str(aluno_folder), "--simular"])
    captured = capsys.readouterr()

    assert exit_code == 0
    assert "simulação" in captured.out
    assert find_aluno_by_name("Marcos anastasio") is None


def test_cli_import_writes_aluno_with_recognized_fields(
    data_dir: Path, aluno_folder: Path, capsys: pytest.CaptureFixture
) -> None:
    exit_code = main(["importar-aluno", str(aluno_folder)])
    capsys.readouterr()

    assert exit_code == 0

    row = _fetch_aluno_row()
    assert row is not None
    assert row["name"] == "Marcos anastasio"
    assert row["objective"] == "definiçao e hipertrofia"
    assert row["plan_start"] == "2024-01-02"
    assert row["plan_end"] == "2024-02-15"


def test_cli_running_twice_does_not_duplicate(
    data_dir: Path, aluno_folder: Path, capsys: pytest.CaptureFixture
) -> None:
    first_exit_code = main(["importar-aluno", str(aluno_folder)])
    capsys.readouterr()

    second_exit_code = main(["importar-aluno", str(aluno_folder)])
    captured = capsys.readouterr()

    assert first_exit_code == 0
    assert second_exit_code == 0
    assert "já existe" in captured.out
    assert _count_alunos() == 1


def test_cli_missing_atleta_label_returns_error(
    data_dir: Path, tmp_path: Path, capsys: pytest.CaptureFixture
) -> None:
    folder = tmp_path / "sem-nome"
    folder.mkdir()
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "objetivo"
    sheet["B1"] = "emagrecimento"
    workbook.save(folder / "treino.xlsx")

    exit_code = main(["importar-aluno", str(folder)])
    captured = capsys.readouterr()

    assert exit_code == 1
    assert "Nome do aluno não encontrado" in captured.out
    # The name check happens before migrations run, so the database is never
    # even created in this scenario.
    assert not db_path().exists()


def test_cli_nonexistent_folder_returns_error_without_raising(
    data_dir: Path, tmp_path: Path, capsys: pytest.CaptureFixture
) -> None:
    exit_code = main(["importar-aluno", str(tmp_path / "does-not-exist")])

    assert exit_code == 1


def test_cli_report_lists_not_imported_items(
    data_dir: Path, aluno_folder: Path, capsys: pytest.CaptureFixture
) -> None:
    main(["importar-aluno", str(aluno_folder), "--simular"])
    captured = capsys.readouterr()

    assert "Não importado" in captured.out


# --- validation against the real Marcos Anastasio folder --------------------


@pytest.mark.skipif(
    not REAL_ALUNO_FOLDER.is_dir(),
    reason="Real validation folder not present on this machine.",
)
def test_parser_against_real_marcos_anastasio_folder() -> None:
    parsed = parse_aluno_folder(REAL_ALUNO_FOLDER)

    assert parsed.fields["name"] == "Marcos anastasio"
    assert parsed.fields["plan_start"] == "2024-01-02"
    assert parsed.fields["plan_end"] == "2024-02-15"
